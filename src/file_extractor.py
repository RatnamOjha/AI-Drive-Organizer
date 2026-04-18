"""
Multi-format file content extractor.
Supports PDF, DOCX, XLSX, TXT, Google Docs/Sheets, and images (via OCR).
"""

from __future__ import annotations

import asyncio
import io
import logging
from typing import Optional

from googleapiclient.http import MediaIoBaseDownload

from .config import Config

logger = logging.getLogger(__name__)

MIME_GOOGLE_DOC = "application/vnd.google-apps.document"
MIME_GOOGLE_SHEET = "application/vnd.google-apps.spreadsheet"
MIME_GOOGLE_SLIDE = "application/vnd.google-apps.presentation"
MIME_PDF = "application/pdf"
MIME_DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
MIME_XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MIME_TXT = "text/plain"
IMAGE_MIMES = {"image/jpeg", "image/png", "image/gif", "image/webp", "image/bmp"}


class FileExtractor:
    """
    Extracts readable text from Google Drive files of various formats.
    Content is truncated to `config.max_content_chars` before classification.
    """

    def __init__(self, config: Config):
        self.config = config
        self._service = None

    def set_service(self, service):
        self._service = service

    async def extract(self, file_id: str, file_name: str, mime_type: str) -> str:
        """Dispatch to the correct extractor based on MIME type."""
        try:
            if mime_type == MIME_GOOGLE_DOC:
                return await self._export_google(file_id, "text/plain")
            elif mime_type == MIME_GOOGLE_SHEET:
                return await self._export_google(file_id, "text/csv")
            elif mime_type == MIME_GOOGLE_SLIDE:
                return await self._export_google(file_id, "text/plain")
            elif mime_type == MIME_PDF:
                return await self._extract_pdf(file_id)
            elif mime_type == MIME_DOCX:
                return await self._extract_docx(file_id)
            elif mime_type == MIME_XLSX:
                return await self._extract_xlsx(file_id)
            elif mime_type == MIME_TXT:
                return await self._download_text(file_id)
            elif mime_type in IMAGE_MIMES and self.config.ocr_enabled:
                return await self._extract_image_ocr(file_id)
            else:
                logger.debug(f"No extractor for MIME type: {mime_type}, using filename only")
                return f"Filename: {file_name}"
        except Exception as e:
            logger.warning(f"⚠️  Extraction failed for '{file_name}': {e}")
            return f"Filename: {file_name}"

    async def _export_google(self, file_id: str, export_mime: str) -> str:
        loop = asyncio.get_event_loop()

        def _export():
            request = self._service.files().export_media(
                fileId=file_id, mimeType=export_mime
            )
            buf = io.BytesIO()
            downloader = MediaIoBaseDownload(buf, request)
            done = False
            while not done:
                _, done = downloader.next_chunk()
            return buf.getvalue().decode("utf-8", errors="replace")

        content = await loop.run_in_executor(None, _export)
        return content[: self.config.max_content_chars]

    async def _download_bytes(self, file_id: str) -> bytes:
        loop = asyncio.get_event_loop()

        def _download():
            request = self._service.files().get_media(fileId=file_id)
            buf = io.BytesIO()
            downloader = MediaIoBaseDownload(buf, request)
            done = False
            while not done:
                _, done = downloader.next_chunk()
            return buf.getvalue()

        return await loop.run_in_executor(None, _download)

    async def _download_text(self, file_id: str) -> str:
        raw = await self._download_bytes(file_id)
        return raw.decode("utf-8", errors="replace")[: self.config.max_content_chars]

    async def _extract_pdf(self, file_id: str) -> str:
        raw = await self._download_bytes(file_id)
        loop = asyncio.get_event_loop()

        def _parse():
            import pdfplumber
            with pdfplumber.open(io.BytesIO(raw)) as pdf:
                pages_text = []
                for page in pdf.pages:
                    t = page.extract_text()
                    if t:
                        pages_text.append(t)
                return "\n".join(pages_text)

        text = await loop.run_in_executor(None, _parse)
        return text[: self.config.max_content_chars]

    async def _extract_docx(self, file_id: str) -> str:
        raw = await self._download_bytes(file_id)
        loop = asyncio.get_event_loop()

        def _parse():
            import docx
            doc = docx.Document(io.BytesIO(raw))
            return "\n".join(p.text for p in doc.paragraphs if p.text.strip())

        text = await loop.run_in_executor(None, _parse)
        return text[: self.config.max_content_chars]

    async def _extract_xlsx(self, file_id: str) -> str:
        raw = await self._download_bytes(file_id)
        loop = asyncio.get_event_loop()

        def _parse():
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            parts = []
            for sheet in wb.worksheets:
                parts.append(f"Sheet: {sheet.title}")
                for row in sheet.iter_rows(max_row=50, values_only=True):
                    row_text = " ".join(str(c) for c in row if c is not None)
                    if row_text.strip():
                        parts.append(row_text)
            return "\n".join(parts)

        text = await loop.run_in_executor(None, _parse)
        return text[: self.config.max_content_chars]

    async def _extract_image_ocr(self, file_id: str) -> str:
        raw = await self._download_bytes(file_id)
        loop = asyncio.get_event_loop()

        def _ocr():
            import pytesseract
            from PIL import Image
            if self.config.tesseract_path:
                pytesseract.pytesseract.tesseract_cmd = self.config.tesseract_path
            image = Image.open(io.BytesIO(raw))
            return pytesseract.image_to_string(image)

        text = await loop.run_in_executor(None, _ocr)
        return text[: self.config.max_content_chars]
